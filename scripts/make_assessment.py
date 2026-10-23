"""生成《项目评估报告》docx + md。数据取自 2026-10-23 实测。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DOCX = Path(r"D:\ai-kb\docs\项目评估报告.docx")
OUT_MD = Path(r"D:\ai-kb\docs\assessment.md")
GREY = RGBColor(0x66, 0x66, 0x66)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
RED = RGBColor(0xB0, 0x2A, 0x2A)

FACTS = [
    ("功能测试", "30 / 30 通过", "tests/run_tests.py，耗时 110 秒"),
    ("检索召回率", "90.0%（27/30）", "30 题测试集，Top-5 命中率"),
    ("答案准确率", "73.3%（22/30）", "LLM 初判；比上次 76.7% 少 1 题（3.3 分，在噪声内）"),
    ("检索质量细节", "平均名次 1.48，Top-1 19/30", "上次为 1.33 / 20"),
    ("检索耗时", "0.92 秒", "上次 0.4 秒；语料从 9,023 涨到 34,395 片"),
    ("问答延迟（热态）", "整轮 6.2 秒；流式首段 2.56 秒", "冷启动首次 22.4 秒"),
    ("入库速率", "2–55 片/秒", "取决于 Ollama 是否被其他程序抢占"),
    ("语料规模", "28 个文件 / 34,395 片 / 735 MB", "其中原 AI 语料仅 9,030 片（26%）"),
    ("代码规模", "后端 3,979 行 / 20 文件；前端 600 行 / 4 文件；测试 439 行", "13 个接口；146 个依赖"),
    ("安全抽查", ".env 未入库、无硬编码密钥、路径穿越已防护", "但接口无鉴权"),
]

SCORES = [
    ("功能完整度", "9.0", "13 个接口全部可用；上传/问答/流式/出处/高亮/历史/模型切换都实测通过"),
    ("回答质量", "7.5", "召回 90%、准确 73.3% 达标；但 Top-1 仅 19/30，无 rerank"),
    ("性能", "7.0", "热态 6.2 秒可接受；检索从 0.4 秒退化到 0.92 秒；入库受争用影响 20 倍"),
    ("健壮性", "6.5", "异常路径覆盖较全；但上传任务在内存、重启即丢；无并发保护"),
    ("安全性", "6.0", "密钥与路径安全做到位；完全没有鉴权，不能直接暴露公网"),
    ("可维护性", "7.5", "模块划分清晰、无 TODO、无裸 except；app 内有 7 处 print 当日志"),
    ("文档", "9.0", "计划书 v1.13、14 篇文档、技术栈/技术点/验收/基线/实验齐全"),
    ("测试", "7.0", "30 个功能用例；缺并发、异常注入与浏览器端到端"),
    ("部署交付", "7.5", "Docker 实机验证、镜像 1.29GB；容器内生成慢 2–4 倍未优化"),
]

FINDINGS = [
    ("P0", "语料领域漂移", 
     "28 个文件里只有 26% 与「AI 学习资料」相关；PostgreSQL 手册（15,669 片）、"
     "Adobe PDF 规范（5,499 片）等占了 74%。检索耗时因此从 0.4 秒涨到 0.92 秒，"
     "产品定位也从「AI 学习知识库」变成了「混合技术库」。",
     "按领域拆集合（ai-kb / tech-ref），检索时指定集合；或移除与目标无关的大部头。"),
    ("P0", "没有任何鉴权",
     "/files、/upload、DELETE /files 全部无需凭证即可调用（实测 HTTP 200）。"
     "本机自用没问题，一旦按 T6.2 部署到公网，任何人都能上传或删除你的资料。",
     "加一个简单 token（请求头 X-Token 比对 .env），或部署侧用 Nginx basic auth。"),
    ("P1", "质量指标粒度太粗",
     "30 题里每题值 3.3 个百分点：本次 73.3% 与上次 76.7% 其实只差 1 道题，"
     "在没有多次采样的情况下无法判断是退化还是噪声。",
     "固定测试集跑 3 次取均值；或扩充到 100 题；并人工抽检 10 题校准 LLM 判定。"),
    ("P1", "入库吞吐受外部因素支配",
     "同一份文件在不同时刻相差 20 倍以上（2 片/秒 vs 55 片/秒）。实测两次是因为"
     "同机另一个项目在调 Ollama，把 bge-m3 挤出显存。",
     "1) 入库前检测 /api/ps，发现别的模型占用就提示；2) 或用独立端口跑第二个 Ollama 实例；"
     "3) 或加任务队列串行化，避免两个上传互相抢显存。"),
    ("P1", "上传任务状态不持久",
     "任务表 _UPLOAD_JOBS 在进程内存里：服务重启后，正在入库的文件状态丢失，"
     "前端轮询会拿到 404，用户不知道到底成没成。",
     "把任务状态写进 SQLite（一行一个任务），重启后可查历史；或启动时扫描 uploads 目录补登记。"),
    ("P1", "检索缺精排",
     "Top-1 命中 19/30，平均名次 1.48：有 11 道题正确片段不在第一位。"
     "计划书的 E2（bge-reranker）一直未做。",
     "对 Top-5 跑一次 rerank 再取前 2 条喂模型；代价是每次多 50–150 毫秒。"),
    ("P2", "前端没有「停止生成」",
     "streamChat 已支持 AbortSignal，但界面上没有取消按钮，长回答只能等。",
     "输入框旁加一个停止按钮，调用 abort()。"),
    ("P2", "日志只有 stdout",
     "7 处 print 代替日志，无请求 ID、无轮转；容器里只能 docker logs 翻。",
     "换 logging + 文件轮转，关键路径带 request_id。"),
    ("P2", "文件列表无分页",
     "28 个文件尚可，上百个会很长；每次刷新拉全量元数据。",
     "加分页或虚拟滚动。"),
    ("P2", "Docker 内生成慢 2–4 倍",
     "容器内问答 13–22 秒，本机 3–6 秒；原因是 Docker Desktop 的 WSL2 与宿主抢 CPU。",
     "真实 Linux 服务器上不明显；若要在 Windows 上演示，录屏用本机模式。"),
]

ROADMAP = [
    ("立刻（半天内）", [
        "给接口加 token 鉴权（改动小，堵住最大的安全洞）",
        "把语料按领域拆成两个集合，问答时指定集合",
    ]),
    ("短期（1–2 天）", [
        "评测跑 3 次取均值，并把结果写进基线文档",
        "上传任务状态落 SQLite，重启可恢复",
        "前端加「停止生成」按钮",
    ]),
    ("中期（1 周）", [
        "接 bge-reranker 做精排（计划书 E2），观察 Top-1 命中率",
        "logging 改造 + 请求 ID",
        "补 3 类测试：并发上传、Ollama 宕机、浏览器端到端",
    ]),
]

PLAN_CHECK = [
    ("F1 切片 512 / 重叠 50", "达标", "E1/E4 实验确定；切片校验 0 超长、重叠 ≥50"),
    ("F2 向量库可持久", "达标", "8,895 → 34,395 片；重启后仍在"),
    ("F3 答案带出处", "达标", "文件名 + 页码 + 相似度 + 片段，可点击回原文并高亮"),
    ("F4 无关问题拒答", "达标", "5/5 拒答（四版提示词对比后定稿）"),
    ("F5 首字 ≤3 秒 / 完整 ≤15 秒", "达标（本机热态）", "流式首段 2.56 秒、整轮 4.3–6.2 秒；冷启动除外"),
    ("F6 召回 ≥80% / 准确 ≥70%", "达标", "90.0% / 73.3%"),
    ("F7 一条命令启动", "达标", "start.bat / docker compose up"),
    ("F8 界面可用", "达标", "多文件上传、进度、模型切换、出处跳转与高亮"),
]


def set_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.25


def add_table(doc, headers, rows, widths, size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(size)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(size)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        cant.set(qn("w:val"), "true")
        tr_pr.append(cant)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)
    doc.add_paragraph()
    return table


def build_docx() -> None:
    doc = Document()
    set_style(doc)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2.2)
        section.left_margin = section.right_margin = Cm(2.3)

    title = doc.add_heading("AI 学习资料知识库 · 项目评估报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("2026-10-23 · 基于代码审查、30 个功能用例与 30 题质量评测的实测结论")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(9)
    sub.runs[0].font.color.rgb = GREY

    doc.add_heading("一、总体结论", level=1)
    p = doc.add_paragraph()
    p.add_run("综合评分 7.5 / 10。").bold = True
    p.add_run(
        "这是一个「功能完整、文档扎实、可演示」的项目：13 个接口全部可用，30 个功能用例全通过，"
        "召回率 90.0%、准确率 73.3% 均达标，Docker 实机验证过。"
        "但它还停留在「个人可用的原型」，离「能对外提供服务」差三件事：没有鉴权、语料领域漂移、"
        "关键指标缺少多次采样。"
    )
    add_table(doc, ["项目", "结论", "依据"], FACTS, [3.4, 5.6, 7.4])

    doc.add_heading("二、分维度评分", level=1)
    add_table(doc, ["维度", "评分", "评价"], SCORES, [3.0, 1.6, 11.8])

    doc.add_heading("三、问题清单（按严重程度）", level=1)
    for level, name, desc, advice in FINDINGS:
        h = doc.add_heading(f"{level} · {name}", level=2)
        if level == "P0":
            for run in h.runs:
                run.font.color.rgb = RED
        doc.add_paragraph(desc)
        p = doc.add_paragraph()
        run = p.add_run("建议：")
        run.bold = True
        p.add_run(advice)

    doc.add_heading("四、与计划书验收目标的对照", level=1)
    add_table(doc, ["验收项", "结论", "依据"], PLAN_CHECK, [4.6, 2.6, 9.2])

    doc.add_heading("五、改进路线", level=1)
    for stage, items in ROADMAP:
        doc.add_heading(stage, level=2)
        for item in items:
            doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("六、值得肯定的地方", level=1)
    for item in [
        "关键结论都有实验支撑：切片长度（E1）、重叠（E4）、上下文条数（E5）、"
        "本地 vs 云端（E6）、拒答提示词（E7）、解析器（E8）——参数不是拍脑袋定的。",
        "踩过的坑都留了记录且修掉：localhost 的 IPv6 回退、langchain embedding 丢 keep_alive、"
        "保活线程覆盖 prompt 缓存、高亮缓存存了坏结果、语料增长触发 SQLite 变量上限。",
        "安全底线守住了：.env 不入库（打包脚本还会主动拦截）、路径穿越已防护、上传有白名单与大小限制。",
        "可交付性完整：一键启动脚本、Docker 编排、四份技术文档、简历与视频脚本。",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    for para in doc.paragraphs:
        if para.style.name.startswith("Heading"):
            para.paragraph_format.keep_with_next = True
    doc.save(OUT_DOCX)


def build_md() -> None:
    def tbl(headers, rows):
        out = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
        out += ["| " + " | ".join(str(x) for x in r) + " |" for r in rows]
        return "\n".join(out)

    md = [
        "# 项目评估报告\n",
        "> 2026-10-23 · 基于代码审查、30 个功能用例与 30 题质量评测的实测结论\n",
        "## 一、总体结论\n",
        "**综合评分 7.5 / 10。** 功能完整、文档扎实、可演示；但离「对外提供服务」差三件事："
        "没有鉴权、语料领域漂移、关键指标缺少多次采样。\n",
        tbl(["项目", "结论", "依据"], FACTS),
        "\n## 二、分维度评分\n",
        tbl(["维度", "评分", "评价"], SCORES),
        "\n## 三、问题清单\n",
    ]
    for level, name, desc, advice in FINDINGS:
        md.append(f"### {level} · {name}\n")
        md.append(desc + "\n")
        md.append(f"**建议**：{advice}\n")
    md += ["\n## 四、与计划书验收目标的对照\n", tbl(["验收项", "结论", "依据"], PLAN_CHECK)]
    md.append("\n## 五、改进路线\n")
    for stage, items in ROADMAP:
        md.append(f"### {stage}\n")
        md += [f"- {i}" for i in items]
        md.append("")
    md.append("\n## 六、值得肯定的地方\n")
    md += [
        "- 关键结论都有实验支撑（E1/E4/E5/E6/E7/E8），参数不是拍脑袋定的。",
        "- 踩过的坑都留了记录且修掉：IPv6 回退、langchain 丢 keep_alive、"
        "保活覆盖 prompt 缓存、高亮坏缓存、SQLite 变量上限。",
        "- 安全底线守住：.env 不入库、路径穿越防护、上传白名单与大小限制。",
        "- 可交付性完整：一键启动、Docker 编排、四份技术文档、简历与视频脚本。",
    ]
    OUT_MD.write_text("\n".join(md), encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_md()
    print("生成:", OUT_DOCX, OUT_DOCX.stat().st_size // 1024, "KB")
    print("生成:", OUT_MD, OUT_MD.stat().st_size // 1024, "KB")
